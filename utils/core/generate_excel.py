import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.table import TableStyleInfo
from pathlib import Path
from PyQt6.QtWidgets import QFileDialog
import utils.logger.logger as log

def generate_excel(data):
    try:
        log.info("Creating Save File Dialog")
        file_path = QFileDialog.getSaveFileName(None, "Save Excel File", "./")
    except Exception as e:
        log.error(f"File Dialog Window Error: {e}")
        return

    if not file_path[0]:
        log.info("User aborted dialog")
        return

    log.info(f"Saving Excel to: {file_path[0]}")

    try:
        dataframes = []
        for subsection, contents in data.items():
            if isinstance(contents, list):
                cleaned = []
                for item in contents:
                    if isinstance(item, dict):
                        flat = {}
                        for key, value in item.items():
                            if not isinstance(value, (list, dict)):
                                flat[key] = value
                        cleaned.append(flat)
                    else:
                        cleaned.append(item)
                contents = cleaned
                table_index = range(len(contents))
            else:
                table_index = [0]
            dataframes.append((subsection, pd.DataFrame(contents, index=table_index)))

        sheet_name = Path(file_path[0]).stem
        vpadding = 2

        with pd.ExcelWriter(file_path[0], engine="openpyxl") as excel_writer:
            current_row = 0
            for name, df in dataframes:
                df.to_excel(excel_writer, sheet_name=sheet_name, startrow=current_row, index=False)
                current_row += len(df) + 1 + vpadding

        excel_file = openpyxl.load_workbook(file_path[0])
        excel_sheet = excel_file[sheet_name]

        # Apply excel table design
        current_row = 1
        table_id = 1
        for name, df in dataframes:
            if df.empty:
                current_row += 1 + vpadding
                continue

            num_rows = len(df)
            num_cols = len(df.columns)
            end_col = get_column_letter(num_cols)
            end_row = current_row + num_rows
            table_ref = f"A{current_row}:{end_col}{end_row}"

            table = Table(displayName=f"Table{table_id}", ref=table_ref)
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleDark8",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False
            )
            excel_sheet.add_table(table)
            table_id += 1
            current_row += num_rows + 1 + vpadding

            # Currency formatting
            currency_format = '"$"#,##0.00'
            currency_columns = {'price', 'total', 'vendor', 'super_x'}

            for row in excel_sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value in currency_columns:
                        col = cell.column
                        for r in range(cell.row + 1, excel_sheet.max_row + 1):
                            target = excel_sheet.cell(row=r, column=col)
                            if target.value is None or isinstance(target.value, str):
                                break
                            if isinstance(target.value, (int, float)):
                                target.number_format = currency_format

            # Auto-width columns
            for column in excel_sheet.columns:
                max_len = 0
                col_letter = column[0].column_letter
                for cell in column:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                excel_sheet.column_dimensions[col_letter].width = max_len + 4
        excel_file.save(file_path[0])

    except Exception as e:
        log.error(f"Excel Creation failed: {e}")